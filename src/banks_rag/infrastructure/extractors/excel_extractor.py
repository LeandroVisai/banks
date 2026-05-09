"""Extracción de chunks desde Excel del Monitor PM.

El archivo del Monitor PM tiene una estructura tabular especial:

- **Filas** = fechas (cada fila es un día de mercado).
- **Columnas** = segmentos narrativos (ej. "Mercado Cambiario", "Renta Fija").
- **Celdas** = texto narrativo completo.

Cada celda con contenido genera un raw_chunk independiente con:
    text = "[YYYY-MM-DD - Nombre Columna]\\n<texto>"
    fecha_iso = fecha de la fila (ISO YYYY-MM-DD)
    page_start/end = número de hoja (1-based)
    section_title_raw = nombre de la columna

Funciones puras: reciben ``Path``, retornan ``(raw_chunks, warnings)``.
"""

from __future__ import annotations

import datetime as _dt
from pathlib import Path
from typing import Any

_MIN_CELL_TEXT_LEN = 10  # celdas con menos de 10 chars se descartan


def find_header_row(rows: list[tuple]) -> tuple[int, list[str]]:
    """Devuelve ``(idx, headers)`` ignorando filas vacías al inicio.

    Una fila se considera header si tiene 2+ celdas no vacías.
    """
    for idx, row in enumerate(rows):
        non_empty = [c for c in row if c is not None and str(c).strip()]
        if len(non_empty) >= 2:
            headers = [str(c).strip() if c is not None else "" for c in row]
            return idx, headers
    return -1, []


def cell_to_str(value: Any) -> str:
    """Convierte un valor de celda a string limpio (datetime → ISO)."""
    if value is None:
        return ""
    if isinstance(value, (_dt.datetime, _dt.date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def extract_cell_chunks(excel_path: Path) -> tuple[list[dict], list[str]]:
    """Extrae chunks a nivel de celda desde un Excel del Monitor PM.

    Retorna ``(raw_chunks, warnings)``. Cada raw_chunk es un dict con keys
    ``text``, ``page_start``, ``page_end``, ``section_title_raw``, ``fecha_iso``.
    """
    warnings: list[str] = []
    raw_chunks: list[dict] = []

    try:
        import openpyxl
    except ImportError:
        warnings.append("missing_openpyxl: pip install openpyxl")
        return raw_chunks, warnings

    try:
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        warnings.append(f"excel_open_error: {e}")
        return raw_chunks, warnings

    for sheet_idx, sheet_name in enumerate(wb.sheetnames, start=1):
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))

        header_idx, headers = find_header_row(rows)
        if header_idx == -1:
            warnings.append(f"sheet_no_header: {sheet_name}")
            continue

        # Detectar columna de fecha: primera con header que contenga "fecha".
        fecha_col: int | None = None
        for ci, h in enumerate(headers):
            if "fecha" in h.lower():
                fecha_col = ci
                break
        if fecha_col is None:
            fecha_col = 0  # fallback

        content_cols = [(ci, h) for ci, h in enumerate(headers) if ci != fecha_col and h]

        for row in rows[header_idx + 1:]:
            if not any(c is not None for c in row):
                continue

            fecha_val = row[fecha_col] if fecha_col < len(row) else None
            fecha_str = cell_to_str(fecha_val)
            if not fecha_str:
                continue

            for ci, col_name in content_cols:
                if ci >= len(row):
                    continue
                cell_text = cell_to_str(row[ci])
                if not cell_text or len(cell_text) < _MIN_CELL_TEXT_LEN:
                    continue

                raw_chunks.append({
                    "text": f"[{fecha_str} - {col_name}]\n{cell_text}",
                    "page_start": sheet_idx,
                    "page_end": sheet_idx,
                    "section_title_raw": col_name,
                    "fecha_iso": fecha_str,
                })

    wb.close()
    return raw_chunks, warnings
